"""Bounded, static imports into canonical steps. Uploaded code is NEVER executed.

Unsupported syntax is an error, not silently discarded. No subprocess, eval,
exec, dependency installation, remote includes, macros or user-defined plugins.
"""
import ast
import csv
import io
import json
import re
import zipfile
from pathlib import PurePosixPath

from app.db import fit_text
from app.step_runtime import MAX_STEPS, normalize_step

MAX_BYTES = 1024 * 1024
MAX_COLUMNS = 40   # see app.step_runtime; vendor exports carry wide metadata
MAX_UNPACKED = 8 * 1024 * 1024
MAX_FILES = 8
# Everything the onboarding understands. app.import_formats adds the market
# formats (plain text, Word, PDF, legacy .xls, Selenium IDE, Postman, Robot,
# Cypress, YAML, Java/C# Selenium, JUnit/TestNG XML) and sniffs the content when
# the extension is unknown, so a team can upload what they already have.
FORMATS = [".csv", ".tsv", ".psv", ".tab", ".xlsx", ".xlsm", ".xls", ".json", ".jsonc", ".py",
           ".js", ".ts", ".mjs", ".jsx", ".tsx", ".feature", ".bdd", ".gherkin", ".robot",
           ".resource", ".java", ".cs", ".side", ".har", ".xml", ".yaml", ".yml", ".txt",
           ".text", ".md", ".markdown", ".log", ".docx", ".doc", ".rtf", ".pdf", ".html",
           ".htm", ".rst", ".step", ".steps", ".jenkins", ".groovy", "Jenkinsfile"]
ALIASES = {"goto": "navigate", "navigate": "navigate", "click": "click", "fill": "fill", "type": "type",
           "press": "press", "check": "check", "uncheck": "uncheck", "select_option": "select_option",
           "selectOption": "select_option", "hover": "hover", "dblclick": "dblclick", "inner_text": "capture_text"}
ASSERTIONS = {"to_have_text": "assert_text", "toHaveText": "assert_text", "to_contain_text": "assert_contains",
              "toContainText": "assert_contains", "to_have_value": "assert_value", "toHaveValue": "assert_value",
              "to_be_visible": "assert_visible", "toBeVisible": "assert_visible", "to_be_hidden": "assert_hidden",
              "toBeHidden": "assert_hidden", "to_be_checked": "assert_checked", "toBeChecked": "assert_checked",
              "to_be_enabled": "assert_enabled", "toBeEnabled": "assert_enabled", "to_be_disabled": "assert_disabled",
              "toBeDisabled": "assert_disabled", "to_have_count": "assert_count", "toHaveCount": "assert_count",
              "to_have_url": "assert_url", "toHaveURL": "assert_url", "to_have_title": "assert_title", "toHaveTitle": "assert_title"}
LOCATORS = {"locator": "css", "get_by_role": "role", "getByRole": "role", "get_by_label": "label", "getByLabel": "label",
            "get_by_text": "text", "getByText": "text", "get_by_test_id": "test_id", "getByTestId": "test_id",
            "get_by_placeholder": "placeholder", "getByPlaceholder": "placeholder"}


class ImportProblem(ValueError):
    pass


class Preview:
    """One import in progress: supported steps, diagnostics and what was detected.

    Errors block saving. Notes and assumptions are recorded and shown but do not
    block, because a real team's spreadsheet always carries metadata columns and
    a precondition line, and refusing the upload over those helps nobody.
    """

    def __init__(self, filename):
        self.filename = filename
        self.steps = []
        self.diagnostics = []
        self.name = PurePosixPath(filename).stem[:200]
        self.start_url = ""
        self.detected = None
        self.execution_plan = None
        self.variables = {}
        # Optional observer every accepted step is reported to (the label grammar
        # uses it to remember which element the flow clicked last). Never required.
        self.on_step = None

    def error(self, line, message):
        self.diagnostics.append({"file": self.filename, "line": line, "severity": "error", "message": str(message)[:400]})

    def note(self, line, message):
        self.diagnostics.append({"file": self.filename, "line": line, "severity": "note", "message": str(message)[:400]})

    def suggest_variables(self, mapping):
        """Variable values found in the file (Robot *** Variables ***, YAML, JSON)."""
        for name, value in (mapping or {}).items():
            key = str(name).strip()
            if re.fullmatch(r"[\w.\-]{1,100}", key):
                self.variables[key] = str(value)

    def add(self, row, line):
        if len(self.steps) >= MAX_STEPS:
            raise ImportProblem("At most 500 steps are allowed")
        try:
            step = normalize_step(row, len(self.steps) + 1)
            step["label"] = fit_text(step.get("label"))
            self.steps.append(step)
            if step["action"] == "navigate" and not self.start_url:
                self.start_url = step["value"]
            if self.on_step is not None:
                try:
                    self.on_step(step, line)
                except Exception:            # an observer must never break an import
                    pass
        except (ValueError, TypeError, OverflowError) as exc:
            self.error(line, exc)

    def result(self):
        errors = [item for item in self.diagnostics if item["severity"] == "error"]
        if not self.steps and not errors:
            self.error(1, "No supported test steps found")
        if sum(s["repeat_count"] for s in self.steps) > 2000:
            self.error(1, "Expanded recording exceeds 2000 actions")
        result = {"filename": self.filename, "name": self.name, "start_url": self.start_url,
                  "steps": self.steps, "diagnostics": self.diagnostics,
                  "can_save": bool(self.steps) and not any(
                      item["severity"] == "error" for item in self.diagnostics),
                  "execution": "sequential", "source_executed": False}
        if self.detected:
            result["detected"] = self.detected
        if self.variables:
            result["variable_suggestions"] = self.variables
        return result


STRICT_COLUMNS = {"action", "selector", "value", "input", "label", "description",
                 "repeat_count", "order", "step", "test_case", "summary", "screenshot"}


def strict_table_header(header):
    """Is this the documented template? (action column, no unknown column)."""
    names = [str(v or "").strip().lower().replace(" ", "_") for v in header]
    return ("action" in names and len(set(names)) == len(names)
            and not (set(names) - STRICT_COLUMNS))


def parse_table(rows, preview):
    """The canonical template: one action per row, optional locator/value columns.

    The Action cell may also carry a step *label* ("Click add to cart",
    "Type: {{password}}") because that is what TestForge's own azure-devops/jira/
    testcomplete exports and other tools write there. app.import_formats owns that
    grammar and the Selector/Input/Expected merging.
    """
    from app.import_formats import import_table_row
    rows = iter(rows)
    header = next(rows, None)
    if not header:
        raise ImportProblem("The table is empty")
    names = [str(v or "").strip().lower().replace(" ", "_") for v in header]
    if "action" not in names:
        raise ImportProblem("Table requires an action column (optional: selector, value, label, repeat_count)")
    if len(set(names)) != len(names):
        raise ImportProblem("Duplicate column names are not supported")
    unknown = set(names) - STRICT_COLUMNS
    if unknown:
        raise ImportProblem("Unknown columns: " + ", ".join(sorted(unknown)))
    case_name = None
    for line, row in enumerate(rows, 2):
        if line > MAX_STEPS + 2:
            raise ImportProblem("At most 500 table rows are allowed")
        if not any(v is not None and str(v).strip() for v in row):
            continue
        if len(row) > len(names) and any(v not in (None, "") for v in row[len(names):]):
            preview.error(line, "Row has more cells than the header")
            continue
        values = dict(zip(names, row))
        case = values.get("test_case") or values.get("summary")
        if case:
            if case_name is not None and str(case) != case_name:
                preview.error(line, "Multiple test cases must be uploaded separately, not flattened into one recording")
            case_name = str(case)
        try:
            import_table_row(preview, {
                "action": values.get("action"),
                "selector": values.get("selector"),
                "value": values.get("value", values.get("input")),
                "label": values.get("description", values.get("label")),
                "repeat": values.get("repeat_count"),
            }, line, "table")
        except (ValueError, TypeError) as exc:
            preview.error(line, exc)


def parse_rows(rows, preview, source="table"):
    """Template first, then the market/vendor headers, then the template error.

    One entry point for CSV/TSV/PSV, .xlsx, .xls, Word tables and sniffed bodies so
    a step table onboards the same way whatever file it arrived in.
    """
    rows = [list(row) for row in rows]
    if strict_table_header(rows[0] if rows else []):
        parse_table(rows, preview)
        return
    from app.import_formats import parse_table_smart
    steps, diagnostics = list(preview.steps), list(preview.diagnostics)
    if not parse_table_smart(rows, preview, source):
        # Nothing recognised the header: report the template's own message rather
        # than "no supported steps found".
        preview.steps, preview.diagnostics = steps, diagnostics
        parse_table(rows, preview)


def parse_excel(data, preview):
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        info = archive.infolist()
        if len(info) > 150 or sum(i.file_size for i in info) > MAX_UNPACKED:
            raise ImportProblem("Excel archive is too large when expanded")
        if any(i.filename.endswith("vbaProject.bin") or "externalLinks/" in i.filename for i in info):
            raise ImportProblem("Macros and external workbook links are not accepted")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ImportProblem("Use one worksheet per upload; multiple sheets must not be silently combined")
        sheet = workbook.worksheets[0]
        if (sheet.max_row or 0) > MAX_STEPS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ImportProblem(f"Excel is limited to {MAX_STEPS} steps and {MAX_COLUMNS} columns")
        rows = []
        for row in sheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                raise ImportProblem("Excel formulas are not test data; replace them with literal values")
            rows.append([cell.value for cell in row])
        parse_rows(rows, preview, "excel")
    finally:
        workbook.close()


class PythonParser:
    def __init__(self, preview):
        self.preview = preview
        self.constants = {}
        self.locators = {}
        self.functions = 0

    def literal(self, node):
        if isinstance(node, ast.Name) and node.id in self.constants:
            return self.constants[node.id]
        try:
            return ast.literal_eval(node)
        except (ValueError, TypeError, SyntaxError) as exc:
            raise ImportProblem("Only literal inputs or literal constants are supported; no evaluated expressions") from exc

    def target(self, node):
        if isinstance(node, ast.Name):
            if node.id == "page":
                return None
            if node.id in self.locators:
                return dict(self.locators[node.id])
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
            method = node.func.attr
            if method in LOCATORS and len(node.args) == 1:
                base = node.func.value
                frames = []
                while isinstance(base, ast.Call) and isinstance(base.func, ast.Attribute) and base.func.attr in {"frame_locator", "frameLocator"}:
                    if len(base.args) != 1 or base.keywords:
                        raise ImportProblem("Frame locator needs one literal selector")
                    frames.insert(0, self.literal(base.args[0]))
                    base = base.func.value
                if not isinstance(base, ast.Name) or base.id != "page":
                    raise ImportProblem("Only page locators and explicit frame chains are supported")
                spec = {"kind": LOCATORS[method], "primary": self.literal(node.args[0])}
                for kw in node.keywords:
                    if kw.arg not in {"name", "exact"}:
                        raise ImportProblem("Unsupported locator option: " + str(kw.arg))
                    spec[kw.arg] = self.literal(kw.value)
                if frames:
                    spec["frames"] = frames
                return spec
        raise ImportProblem("Unsupported locator expression; use page.locator/get_by_role/label/text/test_id")

    def call(self, node, line):
        if isinstance(node, ast.Await):
            node = node.value
        if not isinstance(node, ast.Call) or not isinstance(node.func, ast.Attribute):
            raise ImportProblem("Only supported Playwright actions are imported")
        method, receiver = node.func.attr, node.func.value
        if method in ASSERTIONS:
            if not isinstance(receiver, ast.Call) or not isinstance(receiver.func, ast.Name) or receiver.func.id != "expect" or len(receiver.args) != 1 or receiver.keywords:
                raise ImportProblem("Assertions must use expect(page or locator)")
            target = self.target(receiver.args[0]) or {}
            action = ASSERTIONS[method]
        else:
            if method not in ALIASES:
                raise ImportProblem("Unsupported Playwright method: " + method)
            target = self.target(receiver) or {}
            action = ALIASES[method]
        args = [self.literal(v) for v in node.args]
        # Legacy page.click(selector), page.fill(selector, value), etc.
        if not target and action not in {"navigate", "assert_url", "assert_title"}:
            if not args:
                raise ImportProblem("A selector is required")
            target = {"primary": args.pop(0)}
        options = {kw.arg: self.literal(kw.value) for kw in node.keywords}
        for key in list(options):
            if key == "timeout":
                target["timeout_ms"] = options.pop(key)
            elif key == "wait_until" and action == "navigate" and options[key] == "domcontentloaded":
                options.pop(key)
            elif key == "label" and action == "select_option" and not args:
                args = [options.pop(key)]
                target["by"] = "label"
        if options:
            raise ImportProblem("Unsupported action options: " + ", ".join(str(k) for k in options))
        needs_value = action in {"navigate", "fill", "type", "press", "select_option", "assert_text", "assert_contains", "assert_value", "assert_count", "assert_url", "assert_title"}
        if len(args) != (1 if needs_value else 0):
            raise ImportProblem("Unexpected arguments for " + method)
        if action == "capture_text":
            raise ImportProblem("Use a canonical capture_text step with an explicit variable name")
        self.preview.add({"action": action, "selector": target, "value": args[0] if args else ""}, line)

    def parse(self, source):
        tree = ast.parse(source)
        if sum(1 for _ in ast.walk(tree)) > 15000:
            raise ImportProblem("Python syntax tree exceeds the import limit")
        self.walk(tree.body)

    def walk(self, body):
        for node in body:
            line = getattr(node, "lineno", 1)
            try:
                if isinstance(node, ast.ImportFrom):
                    if node.module not in {"playwright.sync_api", "playwright.async_api"} or any(n.name not in {"expect", "sync_playwright", "async_playwright", "Page"} or n.asname for n in node.names):
                        raise ImportProblem("Only standard Playwright imports are supported")
                elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    self.functions += 1
                    if self.functions > 1 or node.decorator_list or node.name not in {"main", "run"} and not node.name.startswith("test_"):
                        raise ImportProblem("Upload one undecorated test function per file; helpers/fixtures need canonical conversion")
                    if any(a.arg != "page" for a in node.args.args) or node.args.defaults or node.args.vararg or node.args.kwarg or node.args.kwonlyargs or node.args.posonlyargs:
                        raise ImportProblem("Only a page fixture parameter is supported")
                    self.walk(node.body)
                elif isinstance(node, ast.Expr) and isinstance(node.value, ast.Constant) and isinstance(node.value.value, str):
                    pass  # module/function docstring
                elif isinstance(node, (ast.Assign, ast.AnnAssign)):
                    targets = node.targets if isinstance(node, ast.Assign) else [node.target]
                    if len(targets) != 1 or not isinstance(targets[0], ast.Name):
                        raise ImportProblem("Only simple literal or locator assignments are supported")
                    name = targets[0].id
                    try:
                        self.locators[name] = self.target(node.value)
                    except ImportProblem:
                        self.constants[name] = self.literal(node.value)
                elif isinstance(node, ast.Expr):
                    self.call(node.value, line)
                else:
                    raise ImportProblem("Unsupported Python syntax: " + type(node).__name__ + "; convert to explicit steps")
            except (ImportProblem, ValueError, TypeError) as exc:
                self.preview.error(line, exc)


def parse_javascript(source, preview):
    # A deliberately small, line-oriented Playwright subset; never masquerades
    # as a JS engine. No callbacks other than a single standard test wrapper.
    parser = PythonParser(preview)
    wrappers = 0
    opened = False
    for line_no, raw in enumerate(source.splitlines(), 1):
        line = raw.strip()
        if not line or line.startswith("//"):
            continue
        if re.fullmatch(r"import\s*\{\s*(?:test\s*,\s*expect|expect\s*,\s*test|test|expect)\s*\}\s*from\s*['\"]@playwright/test['\"];?", line):
            continue
        if re.fullmatch(r"(?:const|let)\s*\{\s*test\s*,\s*expect\s*\}\s*=\s*require\(['\"]@playwright/test['\"]\);?", line):
            continue
        if re.fullmatch(r"test\(\s*(['\"])[^'\"]+\1\s*,\s*async\s*\(\s*\{\s*page\s*\}\s*\)\s*=>\s*\{", line):
            wrappers += 1
            if wrappers > 1:
                preview.error(line_no, "Upload one test per file; separate tests must not be flattened")
            opened = True
            continue
        if line in {"});", "})"}:
            if not opened:
                preview.error(line_no, "Unexpected test wrapper closure")
            opened = False
            continue
        line = re.sub(r"^await\s+", "", line).removesuffix(";")
        try:
            # JS locator options are restricted to a literal role name/exact.
            line = re.sub(r",\s*\{\s*(name|exact)\s*:\s*([^{}]+)\s*\}\s*\)", lambda m: ", " + m[1] + "=" + ({"true": "True", "false": "False"}.get(m[2].strip(), m[2]) if m[1] == "exact" else m[2]) + ")", line)
            node = ast.parse(line, mode="eval").body
            parser.call(node, line_no)
        except (SyntaxError, ValueError, TypeError, ImportProblem) as exc:
            preview.error(line_no, "Unsupported JS/TS statement: " + str(exc))
    if opened:
        preview.error(len(source.splitlines()), "Unclosed test wrapper")


def parse_bdd(source, preview):
    scenarios = 0
    quoted = r'"((?:[^"\\]|\\.)*)"'
    for line_no, raw in enumerate(source.splitlines(), 1):
        line = raw.strip()
        if not line or line.startswith("#") or line.startswith("Feature:"):
            continue
        if line.startswith("Scenario:"):
            scenarios += 1
            if scenarios > 1:
                preview.error(line_no, "Upload one scenario per file")
            preview.name = line.split(":", 1)[1].strip()[:200]
            continue
        match = re.match(r"^(Given|When|Then|And|But)\s+(.+)$", line)
        if not match:
            preview.error(line_no, "Unsupported Gherkin syntax (outlines, tables, tags and custom step definitions require conversion)")
            continue
        phrase = match[2]
        rules = [
            (r"I (?:open|navigate to) " + quoted, "navigate", False),
            (r"I (click|check|uncheck|hover) " + quoted, "control", False),
            (r"I (fill|type into) " + quoted + r" with " + quoted, "input", False),
            (r"I select " + quoted + r" from " + quoted, "select_option", False),
            (quoted + r" should (contain|equal|have value) " + quoted, "assert", False),
            (quoted + r" should be (visible|hidden|checked|enabled|disabled)", "state", False),
            (r"the URL should be " + quoted, "assert_url", False),
        ]
        found = False
        for pattern, kind, _ in rules:
            m = re.fullmatch(pattern, phrase)
            if not m:
                continue
            found = True
            parts = [json.loads('"' + v + '"') for v in m.groups()]
            row = {"action": kind, "value": "", "selector": {}}
            if kind in {"navigate", "assert_url"}:
                row["value"] = parts[0]
            elif kind == "control":
                row.update(action=parts[0], selector=parts[1])
            elif kind == "input":
                row.update(action="fill" if parts[0] == "fill" else "type", selector=parts[1], value=parts[2])
            elif kind == "select_option":
                row.update(value=parts[0], selector=parts[1])
            elif kind == "assert":
                row.update(action={"contain": "assert_contains", "equal": "assert_text", "have value": "assert_value"}[parts[1]], selector=parts[0], value=parts[2])
            elif kind == "state":
                row.update(action="assert_" + parts[1], selector=parts[0])
            preview.add(row, line_no)
            break
        if not found:
            preview.error(line_no, "Unknown BDD phrase; use a supported phrase or canonical CSV/JSON action")


def _parse_delimited(source, suffix, preview):
    """A delimited table, strict template first, then the market headers."""
    delimiter = "\t" if suffix == ".tsv" else ("|" if suffix == ".psv" else ("," if suffix == ".csv" else "\t"))
    rows = list(csv.reader(io.StringIO(source), delimiter=delimiter))
    parse_rows(rows, preview, "table")
    return True


def parse_file(filename, data):
    preview = Preview(PurePosixPath(filename.replace("\\", "/")).name)
    execution_plan = None
    try:
        if not data or len(data) > MAX_BYTES:
            raise ImportProblem("Upload must be nonempty and at most 1 MiB")
        suffix = PurePosixPath(filename).suffix.lower()
        stem = PurePosixPath(filename).stem.lower()
        # Formats the strict parser does not own: plain text and the market
        # formats, plus content sniffing for an extension nobody listed.
        from app import import_formats
        if import_formats.parse_extra(preview, filename, data) is not None:
            pass
        elif suffix == ".xlsx":
            parse_excel(data, preview)
        elif suffix == ".xlsm":
            parse_excel(data, preview)
        else:
            if stem == "jenkinsfile":
                raise ImportProblem("A Jenkinsfile needs the script it references: upload both files together")
            source = data.decode("utf-8-sig")
            if suffix in {".csv", ".tsv", ".psv", ".tab"}:
                _parse_delimited(source, suffix, preview)
            elif suffix == ".json":
                obj = json.loads(source)
                if isinstance(obj, dict) and 'execution_plan' in obj:
                    from app.plan_runtime import normalize_plan
                    execution_plan = normalize_plan(obj['execution_plan'])
                    steps = execution_plan['steps']
                elif isinstance(obj, dict) and (obj.get('tests') or obj.get('commands')):
                    from app.import_formats import parse_selenium_side
                    parse_selenium_side(preview, obj)
                    steps = None
                elif isinstance(obj, dict) and isinstance(obj.get('item'), list):
                    from app.import_formats import parse_postman
                    parse_postman(preview, obj)
                    steps = None
                else:
                    steps = obj.get("steps") if isinstance(obj, dict) else obj
                if isinstance(obj, dict):
                    preview.name = str(obj.get("name") or preview.name)[:200]
                    variables = obj.get("variables")
                    if isinstance(variables, dict):
                        preview.suggest_variables(variables)
                if steps is not None:
                    if not isinstance(steps, list):
                        raise ImportProblem("JSON must be a step array or an object containing steps")
                    for i, step in enumerate(steps, 1):
                        preview.add(step, i)
            elif suffix == ".py":
                PythonParser(preview).parse(source)
            elif suffix in {".js", ".ts", ".mjs", ".jsx", ".tsx"}:
                # A Cypress spec is JavaScript too: read it with the Cypress
                # grammar instead of reporting every cy.* call as Playwright.
                if import_formats.CYPRESS_HINT.search(source):
                    import_formats.parse_cypress(preview, source)
                    preview.detected = {"format": "js", "why": "Cypress spec (cy.* commands)"}
                else:
                    parse_javascript(source, preview)
            elif suffix in {".feature", ".bdd"}:
                parse_bdd(source, preview)
            elif suffix in {".groovy", ".jenkins"} or preview.filename.lower() == "jenkinsfile":
                raise ImportProblem("Jenkins is an orchestration file, not browser steps. Upload it together with its referenced Python/Playwright script using bundle import; shell commands are never executed.")
            else:
                # An extension nobody listed: read the content and decide.
                detected = import_formats.sniff(source)
                preview.detected = {"format": detected["format"], "why": detected["why"],
                                    "extension": suffix or "(none)"}
                import_formats._parse_by_sniff(preview, source, detected["format"], filename)
    except ImportProblem as exc:
        preview.error(getattr(exc, "lineno", 1), exc)
    except (ValueError, TypeError, SyntaxError, RecursionError, zipfile.BadZipFile, KeyError, OverflowError) as exc:
        preview.error(getattr(exc, "lineno", 1), exc)
    result = preview.result()
    if execution_plan is not None:
        result['execution_plan'] = execution_plan
    elif preview.execution_plan is not None:
        result['execution_plan'] = preview.execution_plan
    if preview.detected:
        result['detected'] = preview.detected
    return result


def parse_bundle(files):
    """One script, optionally wrapped by a restricted declarative Jenkinsfile.

    Reject non-test shell commands instead of giving a misleading green import.
    """
    if not files or len(files) > MAX_FILES or sum(len(data) for _, data in files) > MAX_BYTES:
        raise ImportProblem("Bundle must contain 1–8 files totalling at most 1 MiB")
    names = [PurePosixPath(name.replace("\\", "/")).name for name, _ in files]
    if len(set(names)) != len(names):
        raise ImportProblem("Duplicate filenames in bundle")
    mapping = dict(zip(names, [data for _, data in files]))
    wrappers = [n for n in names if n.lower() == "jenkinsfile" or PurePosixPath(n).suffix.lower() in {".jenkins", ".groovy"}]
    if not wrappers:
        if len(files) != 1:
            raise ImportProblem("Upload one test script per recording, or a Jenkinsfile with one referenced script")
        return parse_file(names[0], files[0][1])
    if len(wrappers) != 1:
        raise ImportProblem("Only one Jenkinsfile per bundle")
    source = mapping[wrappers[0]].decode("utf-8-sig")
    commands = re.findall(r"\bsh\s*\(?\s*(['\"])(.*?)\1\s*\)?", source)
    remainder = re.sub(r"\bsh\s*\(?\s*(['\"])(.*?)\1\s*\)?", "", source)
    remainder = re.sub(r"//[^\n]*", "", remainder)
    remainder = re.sub(r"\bstage\s*\(\s*(['\"])[^'\"]*\1\s*\)", "", remainder)
    remainder = re.sub(r"\b(?:pipeline|stages|steps|agent\s+any)\b|[{};\s]", "", remainder)
    if remainder or len(commands) != 1:
        raise ImportProblem("Only a declarative Jenkinsfile with agent any, stages/steps and ONE literal test command is supported; no checkout, installers, credentials, post hooks or other shell code")
    command = commands[0][1]
    match = re.fullmatch(r"(?:python3?|pytest|node|npx playwright test)\s+([A-Za-z0-9_.-]+\.(?:py|js|ts))", command)
    if not match or match[1] not in mapping:
        raise ImportProblem("Upload the referenced script; allowed Jenkins command: python/pytest/node/npx playwright test <filename>")
    if set(mapping) != {wrappers[0], match[1]}:
        raise ImportProblem("Unreferenced files would be ignored; remove them from this bundle")
    result = parse_file(match[1], mapping[match[1]])
    result["wrapper"] = wrappers[0]
    result["notice"] = "Jenkins wrapper statically validated. Only its referenced test steps are imported; no pipeline or shell command runs."
    return result
